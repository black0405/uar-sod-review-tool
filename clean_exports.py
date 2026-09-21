#!/usr/bin/env python3
"""Standalone cleaner: drop blocked + sanctioned rows from apps & websites exports.

Put everything in one folder:
    apps export(s)      .xlsx / .csv   (any name)
    websites export(s)  .xlsx / .csv   (any name)
    blocked list(s)     .xlsx / .csv   file name OR sheet name contains "block"

Run:   python clean_exports.py <folder>          (folder defaults to cwd)
Out:   <folder>/cleaned/<same file name>  — same columns, blocked/sanctioned rows removed

Matching mirrors unsanctioned-review.html: name and domain compared lower-cased,
trimmed, with surrounding [brackets] stripped. Needs: pip install openpyxl
"""
import csv
import re
import sys
from pathlib import Path

import openpyxl

# column aliases -> canonical field (same table as the HTML tool)
COLMAP = {
    "application": "name", "applications": "name", "app": "name",
    "website": "name", "websites": "name", "site": "name",
    "domain": "domain", "domains": "domain",
    "sanctioned": "sanctioned",
    "category": "category", "appcategory": "category",
    "cci": "cci", "ccl": "cci",
    "users": "users", "user": "users", "numberofusers": "users",
    "bytesuploaded": "up", "uploadbytes": "up", "bytesupload": "up",
    "bytesdownloaded": "down", "downloadbytes": "down",
    "totalbytes": "total", "bytestotal": "total",
    "sessions": "sessions", "session": "sessions",
    "pagevisit": "visits", "pagevisits": "visits", "pagecount": "visits",
    "tags": "tags", "tag": "tags",
}
FUZZY = [
    (r"categor", "category"), (r"application|appname", "name"),
    (r"website|sitename|^url$", "name"), (r"domain", "domain"),
    (r"upload|sent", "up"), (r"download|receive|recv", "down"),
    (r"user", "users"), (r"session", "sessions"), (r"cci|confidence", "cci"),
    (r"sanction", "sanctioned"), (r"visit|page", "visits"),
]
# cells in a blocked sheet that are headers/noise, not names
NOISE = re.compile(r"^(zone|category|cci|users?|description|.*bytes.*|.*traffic.*|.*comment.*|.*action.*|application|website|domain|sanctioned)$")


def norm(v) -> str:
    return "" if v is None else str(v).strip().lower()


def strip_br(v) -> str:
    return re.sub(r"^\[+\s*|\s*\]+$", "", norm(v))


def map_header(h):
    k = re.sub(r"[#_\s]+", "", norm(h))
    if not k:
        return None
    if k in COLMAP:
        return COLMAP[k]
    if re.search(r"total", k) and re.search(r"byte|traffic", k):
        return "total"
    return next((f for pat, f in FUZZY if re.search(pat, k)), None)


def find_header(rows):
    """First row (of the first 30) with a name column and >=2 known columns.
    Returns (index, field-per-column) or (None, None)."""
    alt = None
    for i, row in enumerate(rows[:30]):
        m = [map_header(c) for c in row]
        n = sum(1 for f in m if f)
        if "name" in m and n >= 2:
            return i, m
        if alt is None and n >= 3:
            alt = (i, m)
    if alt:
        i, m = alt
        for j, h in enumerate(rows[i]):   # first unmapped non-empty header = name
            if not m[j] and norm(h):
                m[j] = "name"
                return i, m
    return None, None


def is_sanctioned(v) -> bool:
    return bool(re.match(r"^(y|true|1|sanction)", norm(v)))


# ── readers / writers: every sheet is a list of row-lists ────────────────────
def read_sheets(path: Path):
    if path.suffix.lower() == ".csv":
        txt = path.read_text(encoding="utf-8-sig", errors="replace")
        head = txt.split("\n", 1)[0]
        delim = max(",;\t|", key=head.count)
        return {"Sheet1": list(csv.reader(txt.splitlines(), delimiter=delim))}, delim
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}, ","


def write_sheets(path: Path, sheets: dict, delim: str):
    if path.suffix.lower() == ".csv":
        with path.open("w", newline="", encoding="utf-8") as f:
            csv.writer(f, delimiter=delim).writerows(next(iter(sheets.values())))
        return
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title[:31])
        for r in rows:
            ws.append(r)
    wb.save(path)


# ── blocked terms ────────────────────────────────────────────────────────────
def blocked_terms(rows) -> set:
    hi, m = find_header(rows)
    out = set()
    if hi is not None:
        ni, di = m.index("name"), (m.index("domain") if "domain" in m else -1)
        for r in rows[hi + 1:]:
            if ni < len(r) and strip_br(r[ni]):
                out.add(strip_br(r[ni]))
            if di >= 0 and di < len(r) and strip_br(r[di]):
                out.add(strip_br(r[di]))
        return out
    for r in rows:   # no header: every text cell is a name
        for c in r:
            s = strip_br(c)
            if len(s) > 1 and not NOISE.match(s) and not re.match(r"^-?\d", s):
                out.add(s)
    return out


# ── clean one data sheet ─────────────────────────────────────────────────────
def clean_rows(rows, blocked: set):
    """Returns (kept rows incl. everything above the header, dropped count) or (None, 0)."""
    hi, m = find_header(rows)
    if hi is None:
        return None, 0
    ni = m.index("name")
    di = m.index("domain") if "domain" in m else -1
    si = m.index("sanctioned") if "sanctioned" in m else -1
    cell = lambda r, i: r[i] if 0 <= i < len(r) else None
    kept, dropped = rows[: hi + 1], 0
    for r in rows[hi + 1:]:
        name, dom = strip_br(cell(r, ni)), strip_br(cell(r, di))
        if not name:
            kept.append(r)
            continue
        if name in blocked or (dom and dom in blocked) or (si >= 0 and is_sanctioned(cell(r, si))):
            dropped += 1
        else:
            kept.append(r)
    return kept, dropped


def main(folder: Path):
    files = sorted(p for p in folder.iterdir() if p.suffix.lower() in (".xlsx", ".xlsm", ".csv") and not p.name.startswith("~$"))
    data, blocked = {}, set()
    for p in files:
        sheets, delim = read_sheets(p)
        blk = {t: r for t, r in sheets.items() if "block" in p.name.lower() or "block" in t.lower()}
        if blk:   # a workbook with any 'block' sheet is a blocked list, not data
            for r in blk.values():
                blocked |= blocked_terms(r)
            print(f"blocked list  {p.name}: {len(blk)} sheet(s)")
            continue
        data[p] = (sheets, delim)
    print(f"{len(blocked)} blocked terms")
    if not data:
        sys.exit("no export files found")
    out_dir = folder / "cleaned"
    out_dir.mkdir(exist_ok=True)
    for p, (sheets, delim) in data.items():
        cleaned, total = {}, 0
        for title, rows in sheets.items():
            kept, dropped = clean_rows(rows, blocked)
            if kept is None:
                continue          # no recognizable columns — sheet skipped
            cleaned[title] = kept
            total += dropped
        if not cleaned:
            print(f"skip          {p.name}: no recognizable columns")
            continue
        write_sheets(out_dir / p.name, cleaned, delim)
        print(f"cleaned       {p.name}: removed {total} rows -> cleaned/{p.name}")


if __name__ == "__main__":
    main(Path(sys.argv[1] if len(sys.argv) > 1 else ".").resolve())
